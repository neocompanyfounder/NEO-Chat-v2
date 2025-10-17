"""File processor service for extracting text from various document formats.

This module provides text extraction from PDF, DOCX, XLSX, PPTX, and TXT files.
"""

import io
from typing import Optional, Dict, Any
from pathlib import Path

# Document processing libraries
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

from src.models.document import DocumentType
from src.utils.logger import get_logger

logger = get_logger(__name__)


class FileProcessorError(Exception):
    """Exception raised for file processing errors."""
    pass


class FileProcessor:
    """Service for extracting text from various document formats."""
    
    def __init__(self):
        """Initialize file processor."""
        self._validate_dependencies()
    
    def _validate_dependencies(self):
        """Validate that required libraries are installed."""
        missing = []
        
        if PyPDF2 is None:
            missing.append("PyPDF2")
        if DocxDocument is None:
            missing.append("python-docx")
        if openpyxl is None:
            missing.append("openpyxl")
        if Presentation is None:
            missing.append("python-pptx")
        
        if missing:
            logger.warning(
                f"Missing optional dependencies: {', '.join(missing)}. "
                "Some file types may not be supported."
            )
    
    async def extract_text(
        self,
        file_content: bytes,
        file_type: DocumentType,
        filename: str
    ) -> Dict[str, Any]:
        """Extract text from a file.
        
        Args:
            file_content: Raw file bytes
            file_type: Type of document
            filename: Original filename
            
        Returns:
            Dictionary with extracted text and metadata
            
        Raises:
            FileProcessorError: If extraction fails
        """
        try:
            if file_type == DocumentType.PDF:
                return await self._extract_from_pdf(file_content, filename)
            elif file_type == DocumentType.DOCX:
                return await self._extract_from_docx(file_content, filename)
            elif file_type == DocumentType.XLSX:
                return await self._extract_from_xlsx(file_content, filename)
            elif file_type == DocumentType.PPTX:
                return await self._extract_from_pptx(file_content, filename)
            elif file_type == DocumentType.TXT:
                return await self._extract_from_txt(file_content, filename)
            else:
                raise FileProcessorError(f"Unsupported file type: {file_type}")
                
        except Exception as e:
            logger.error(
                f"Failed to extract text from {filename}: {str(e)}",
                extra={
                    "filename": filename,
                    "file_type": file_type.value,
                    "error": str(e)
                }
            )
            raise FileProcessorError(f"Text extraction failed: {str(e)}")
    
    async def _extract_from_pdf(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text from PDF file.
        
        Args:
            content: PDF file bytes
            filename: Original filename
            
        Returns:
            Extracted text and metadata
        """
        if PyPDF2 is None:
            raise FileProcessorError("PyPDF2 not installed")
        
        try:
            pdf_file = io.BytesIO(content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text_parts = []
            page_count = len(pdf_reader.pages)
            
            for page_num, page in enumerate(pdf_reader.pages, 1):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"[Page {page_num}]\n{page_text}")
                except Exception as e:
                    logger.warning(
                        f"Failed to extract text from page {page_num}: {str(e)}",
                        extra={"filename": filename, "page": page_num}
                    )
            
            full_text = "\n\n".join(text_parts)
            
            # Extract metadata
            metadata = {
                "pages": page_count,
                "format": "pdf"
            }
            
            if pdf_reader.metadata:
                if pdf_reader.metadata.title:
                    metadata["title"] = pdf_reader.metadata.title
                if pdf_reader.metadata.author:
                    metadata["author"] = pdf_reader.metadata.author
            
            logger.info(
                f"Extracted text from PDF: {page_count} pages",
                extra={
                    "filename": filename,
                    "pages": page_count,
                    "text_length": len(full_text)
                }
            )
            
            return {
                "text": full_text,
                "metadata": metadata
            }
            
        except Exception as e:
            raise FileProcessorError(f"PDF extraction failed: {str(e)}")
    
    async def _extract_from_docx(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text from DOCX file.
        
        Args:
            content: DOCX file bytes
            filename: Original filename
            
        Returns:
            Extracted text and metadata
        """
        if DocxDocument is None:
            raise FileProcessorError("python-docx not installed")
        
        try:
            docx_file = io.BytesIO(content)
            doc = DocxDocument(docx_file)
            
            # Extract paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            full_text = "\n\n".join(paragraphs)
            
            # Extract tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        table_texts.append(row_text)
            
            if table_texts:
                full_text += "\n\n[Tables]\n" + "\n".join(table_texts)
            
            metadata = {
                "paragraphs": len(paragraphs),
                "tables": len(doc.tables),
                "format": "docx"
            }
            
            logger.info(
                f"Extracted text from DOCX: {len(paragraphs)} paragraphs",
                extra={
                    "filename": filename,
                    "paragraphs": len(paragraphs),
                    "text_length": len(full_text)
                }
            )
            
            return {
                "text": full_text,
                "metadata": metadata
            }
            
        except Exception as e:
            raise FileProcessorError(f"DOCX extraction failed: {str(e)}")
    
    async def _extract_from_xlsx(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text from XLSX file.
        
        Args:
            content: XLSX file bytes
            filename: Original filename
            
        Returns:
            Extracted text and metadata
        """
        if openpyxl is None:
            raise FileProcessorError("openpyxl not installed")
        
        try:
            xlsx_file = io.BytesIO(content)
            workbook = openpyxl.load_workbook(xlsx_file, data_only=True)
            
            sheet_texts = []
            total_rows = 0
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_texts.append(f"[Sheet: {sheet_name}]")
                
                for row in sheet.iter_rows(values_only=True):
                    # Filter out None values and convert to strings
                    row_values = [str(cell) for cell in row if cell is not None]
                    if row_values:
                        sheet_texts.append(" | ".join(row_values))
                        total_rows += 1
            
            full_text = "\n".join(sheet_texts)
            
            metadata = {
                "sheets": len(workbook.sheetnames),
                "rows": total_rows,
                "format": "xlsx"
            }
            
            logger.info(
                f"Extracted text from XLSX: {len(workbook.sheetnames)} sheets",
                extra={
                    "filename": filename,
                    "sheets": len(workbook.sheetnames),
                    "rows": total_rows
                }
            )
            
            return {
                "text": full_text,
                "metadata": metadata
            }
            
        except Exception as e:
            raise FileProcessorError(f"XLSX extraction failed: {str(e)}")
    
    async def _extract_from_pptx(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text from PPTX file.
        
        Args:
            content: PPTX file bytes
            filename: Original filename
            
        Returns:
            Extracted text and metadata
        """
        if Presentation is None:
            raise FileProcessorError("python-pptx not installed")
        
        try:
            pptx_file = io.BytesIO(content)
            prs = Presentation(pptx_file)
            
            slide_texts = []
            
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_texts.append(f"[Slide {slide_num}]")
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text)
            
            full_text = "\n\n".join(slide_texts)
            
            metadata = {
                "slides": len(prs.slides),
                "format": "pptx"
            }
            
            logger.info(
                f"Extracted text from PPTX: {len(prs.slides)} slides",
                extra={
                    "filename": filename,
                    "slides": len(prs.slides),
                    "text_length": len(full_text)
                }
            )
            
            return {
                "text": full_text,
                "metadata": metadata
            }
            
        except Exception as e:
            raise FileProcessorError(f"PPTX extraction failed: {str(e)}")
    
    async def _extract_from_txt(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text from TXT file.
        
        Args:
            content: TXT file bytes
            filename: Original filename
            
        Returns:
            Extracted text and metadata
        """
        try:
            # Try UTF-8 first, fall back to latin-1
            try:
                text = content.decode('utf-8')
            except UnicodeDecodeError:
                text = content.decode('latin-1')
            
            metadata = {
                "lines": len(text.split('\n')),
                "format": "txt"
            }
            
            logger.info(
                f"Extracted text from TXT: {len(text)} characters",
                extra={
                    "filename": filename,
                    "text_length": len(text)
                }
            )
            
            return {
                "text": text,
                "metadata": metadata
            }
            
        except Exception as e:
            raise FileProcessorError(f"TXT extraction failed: {str(e)}")
    
    def detect_file_type(self, filename: str, mime_type: str) -> DocumentType:
        """Detect file type from filename and MIME type.
        
        Args:
            filename: Original filename
            mime_type: MIME type
            
        Returns:
            Detected document type
        """
        # Check file extension
        ext = Path(filename).suffix.lower()
        
        if ext == '.pdf' or 'pdf' in mime_type:
            return DocumentType.PDF
        elif ext == '.docx' or 'wordprocessingml' in mime_type:
            return DocumentType.DOCX
        elif ext == '.xlsx' or 'spreadsheetml' in mime_type:
            return DocumentType.XLSX
        elif ext == '.pptx' or 'presentationml' in mime_type:
            return DocumentType.PPTX
        elif ext == '.txt' or 'text/plain' in mime_type:
            return DocumentType.TXT
        elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp'] or 'image/' in mime_type:
            return DocumentType.IMAGE
        else:
            return DocumentType.UNKNOWN
